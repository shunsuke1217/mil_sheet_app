"""xlsx を PC 内に保存する（SPEC 12.2）。重量はアプリで計算せず、Excel の数式で出す（8.5 節）。"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from millsheet.export.base import save_atomic

COLUMNS = [  # (見出し, レコードのキー)。重量はキーを持たず、数式を入れる
    ("ページ", "ページ"), ("材種", "材種"), ("材質", "材質"), ("寸法", "寸法"), ("長さ(mm)", "長さ_mm"),
    ("本数", "本数"), ("重量(kg)", None), ("要確認", "要確認"), ("確認理由", "確認理由"),
    ("元本数", "元本数"), ("質量(kg)", "質量"),
]
YELLOW = PatternFill("solid", fgColor="FFF9C4")
RED = PatternFill("solid", fgColor="FFCDD2")
WEIGHT_COLUMN = 7
NUMBER_FORMAT_COLUMNS = {5, WEIGHT_COLUMN}
VALUE_COLUMNS = {1, 2, 3, 4, 5, 6, WEIGHT_COLUMN, 10, 11}   # 空なら赤くする列（要確認・確認理由の列は除く）
PL_BLANK_COLUMNS = {5}             # PL は長さが空欄で正しい


def weight_formula(row: int) -> str:
    """重量 = 質量 ÷ 元本数 × 本数（四捨五入）。本数が元本数以上なら質量そのまま。どれかが空なら空。"""
    qty, printed, mass = f"F{row}", f"J{row}", f"K{row}"
    return (f'=IF(OR({qty}="",{printed}="",{mass}="",{printed}=0),"",'
            f'ROUND(IF({qty}>={printed},{mass},{mass}/{printed}*{qty}),0))')


def _weight_is_blank(rec: dict) -> bool:
    return rec.get("本数") is None or not rec.get("元本数") or rec.get("質量") is None


class LocalXlsxExporter:
    def __init__(self, path: Path):
        self.path = Path(path)

    def export(self, records: list[dict], meta: dict) -> str:
        wb = Workbook()
        ws = wb.active
        ws.title = "ミルシート"
        ws.append([h for h, _ in COLUMNS])
        for c in ws[1]:
            c.font = Font(bold=True)
        for row_no, rec in enumerate(records, start=2):
            for col, (_, key) in enumerate(COLUMNS, start=1):
                if col == WEIGHT_COLUMN:
                    v = weight_formula(row_no)
                    blank = _weight_is_blank(rec)
                else:
                    v = rec.get(key)
                    if key == "要確認":
                        v = "要確認" if v else None
                    elif key == "確認理由":
                        v = "、".join(v) if v else None
                    blank = v is None
                cell = ws.cell(row_no, col, v)
                if col in NUMBER_FORMAT_COLUMNS:
                    cell.number_format = "#,##0"
                pl_blank = rec.get("材種") == "PL" and col in PL_BLANK_COLUMNS
                if blank and col in VALUE_COLUMNS and not pl_blank:
                    cell.fill = RED
                elif rec.get("要確認"):
                    cell.fill = YELLOW
        last = 1 + len(records)
        ws.cell(last + 1, WEIGHT_COLUMN - 1, "合計").font = Font(bold=True)
        total = ws.cell(last + 1, WEIGHT_COLUMN, f"=SUM(G2:G{last})")
        total.number_format = "#,##0"
        total.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:K{last}"
        for col, width in zip("ABCDEFGHIJK", (8, 12, 10, 22, 10, 7, 10, 8, 60, 8, 10)):
            ws.column_dimensions[col].width = width
        save_atomic(self.path, lambda p: wb.save(p))
        return str(self.path)
