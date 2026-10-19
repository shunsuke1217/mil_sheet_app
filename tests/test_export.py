import json
import os
from pathlib import Path

import openpyxl
import pytest

from millsheet.export import get_exporter
from millsheet.export.base import OutputLockedError
from millsheet.export.json_out import write_json
from millsheet.export.local_xlsx import LocalXlsxExporter
from millsheet.rawfile import read_raw, write_raw

HEADER = ["ページ", "材種", "材質", "寸法", "長さ(mm)", "本数", "重量(kg)", "要確認", "確認理由", "元本数", "質量(kg)"]


def weight_formula(r):
    return f'=IF(OR(F{r}="",J{r}="",K{r}="",J{r}=0),"",ROUND(IF(F{r}>=J{r},K{r},K{r}/J{r}*F{r}),0))'


def _rgb(cell):
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None


def test_xlsx_layout(tmp_path, ordered_sample_dicts):
    p = LocalXlsxExporter(tmp_path / "out.xlsx").export(ordered_sample_dicts, {})
    assert p == str(tmp_path / "out.xlsx")
    ws = openpyxl.load_workbook(p)["ミルシート"]
    assert [c.value for c in ws[1]] == HEADER
    # 並べると例外サンプルの Ｈ-294×200×8×12（12000, 10000）が 2・3 行目、サンプル _1 が 4 行目に来る
    assert ws["A2"].value == "Ｐ-1" and ws["A4"].value == "Ｐ-2"
    # 重量はアプリで計算せず、同じ行の本数・元本数・質量から Excel が計算する（人が質量を打ち込めば重量が出る）
    assert (ws["F4"].value, ws["J4"].value, ws["K4"].value) == (1, 18, 10332)
    assert ws["G4"].value == weight_formula(4) and ws["G2"].value == weight_formula(2)
    assert ws["G4"].number_format == "#,##0" and ws["E4"].number_format == "#,##0"
    last = 1 + len(ordered_sample_dicts)
    assert ws.cell(last + 1, 7).value == f"=SUM(G2:G{last})" and ws.cell(last + 1, 6).value == "合計"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == f"A1:K{last}"


def test_xlsx_colors(tmp_path, ordered_sample_dicts):
    ws = openpyxl.load_workbook(LocalXlsxExporter(tmp_path / "o.xlsx").export(ordered_sample_dicts, {}))["ミルシート"]
    review_row = next(r for r in range(2, ws.max_row) if (ws.cell(r, 9).value or "").startswith("丸数字が書かれていません"))
    assert ws.cell(review_row, 8).value == "要確認"
    assert _rgb(ws.cell(review_row, 2)).endswith("FFF9C4")     # 品名あり → 黄色
    assert _rgb(ws.cell(review_row, 6)).endswith("FFCDD2")     # 本数 null → 赤
    assert _rgb(ws.cell(review_row, 7)).endswith("FFCDD2")     # 本数が無いので重量も出ない → 赤
    assert _rgb(ws.cell(review_row, 11)).endswith("FFF9C4")    # 質量は読めている → 黄色
    assert _rgb(ws.cell(4, 2)) is None                         # 要確認でない行は塗らない
    pl_row = next(r for r in range(2, ws.max_row) if ws.cell(r, 2).value == "PL")
    assert ws.cell(pl_row, 5).value is None and _rgb(ws.cell(pl_row, 5)) is None


def test_xlsx_locked_raises(tmp_path, monkeypatch, ordered_sample_dicts):
    def deny(*a):
        raise PermissionError()
    monkeypatch.setattr(os, "replace", deny)
    with pytest.raises(OutputLockedError):
        LocalXlsxExporter(tmp_path / "o.xlsx").export(ordered_sample_dicts, {})


def test_get_exporter(tmp_path):
    assert isinstance(get_exporter("local_xlsx", tmp_path / "o.xlsx"), LocalXlsxExporter)
    with pytest.raises(ValueError):
        get_exporter("google_sheets", tmp_path / "o.xlsx")


def test_write_json(tmp_path, ordered_sample_dicts):
    write_json(tmp_path / "o.json", ordered_sample_dicts, {"モデル": "m"})
    data = json.loads((tmp_path / "o.json").read_text("utf-8"))
    rec = data["records"][2]
    assert data["meta"] == {"モデル": "m"} and (rec["元本数"], rec["質量"]) == (18, 10332) and "重量" not in rec


def test_raw_roundtrip(tmp_path, sample_pages):
    write_raw(tmp_path / "x_raw.json", tmp_path, sample_pages)
    src, pages = read_raw(tmp_path / "x_raw.json")
    assert src == tmp_path and pages == sample_pages


def test_read_raw_accepts_fixture():
    src, pages = read_raw(Path("tests/fixtures/sample_raw.json"))
    assert src is None and len(pages) == 12
    assert pages[0].file == "20261014084620.pdf" and pages[-1].error == "APIタイムアウト"
